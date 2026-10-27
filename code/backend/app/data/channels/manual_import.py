"""人工导入（07 号文档 6.6 节）：按表头特征匹配解析器（不依赖文件名）→ 逐行校验 → 全部通过才入库。

每种导出模板一个解析器。表头以 POC 时拿到的真实导出文件为准，模板变了在这里加新版本解析器。
"""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.adapters import types as T
from app.data import store
from app.data.items import load_items
from app.db import utcnow
from app.models import FetchJob, ManualImport, Shop
from app.workflow.events import emit


def _int(v: str) -> int:
    return int(str(v).replace(",", "").strip())


def _fen(v: str) -> int:
    """「239.20」元 → 23920 分。"""
    try:
        return int((Decimal(str(v).replace(",", "").replace("¥", "").strip()) * 100).quantize(Decimal(1)))
    except InvalidOperation as e:
        raise ValueError(f"不是金额：{v}") from e


def _day(v: str) -> date:
    return date.fromisoformat(str(v).strip().replace("/", "-")[:10])


def _str(v: str) -> str:
    return str(v).strip()


@dataclass(frozen=True)
class Parser:
    name: str
    item: str
    # 表头 → (字段名, 转换函数)；required 全部出现才算匹配，optional 有就读
    required: dict[str, tuple[str, Callable[[str], Any]]]
    optional: dict[str, tuple[str, Callable[[str], Any]]] = field(default_factory=dict)

    def matches(self, header: list[str]) -> bool:
        return set(self.required) <= set(header)


PARSERS = [
    Parser(
        "pdd_goods_traffic_v1", "goods_daily_traffic",
        required={"日期": ("day", _day), "商品ID": ("goods_id", _str), "商品访客数": ("visitors", _int),
                  "商品加购件数": ("add_to_cart", _int), "支付订单数": ("orders", _int)},
        optional={"支付金额": ("gmv_fen", _fen), "退款订单数": ("refunds", _int)},
    ),
    Parser(
        "pdd_order_stats_v1", "order_stats",
        required={"日期": ("day", _day), "成交订单数": ("orders", _int), "成交金额": ("gmv_fen", _fen)},
        optional={"退款订单数": ("refunds", _int)},
    ),
]


def _read_rows(path: Path) -> list[list[str]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True, data_only=True).active
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        return []
    return [row for row in csv.reader(io.StringIO(text)) if any(c.strip() for c in row)]


def import_file(s: Session, shop: Shop, path: Path, uploaded_by: str) -> ManualImport:
    rec = ManualImport(file_path=str(path), uploaded_by=uploaded_by, rows=0)
    s.add(rec)
    try:
        rows = _read_rows(Path(path))
    except Exception:  # noqa: BLE001  打不开的文件按「不认识」处理
        rows = []
    header = [h.strip() for h in rows[0]] if rows else []
    parser = next((p for p in PARSERS if p.matches(header)), None)
    if parser is None:
        rec.status = "unrecognized"
        s.flush()
        return rec
    rec.parser, rec.item = parser.name, parser.item
    idef = load_items()[parser.item]
    schema = getattr(T, idef.schema)
    cols = {**parser.optional, **parser.required}
    latest_ok = (utcnow() + timedelta(hours=8)).date()

    by_day: dict[date, list] = defaultdict(list)
    errors: list[dict[str, Any]] = []
    for line_no, row in enumerate(rows[1:], start=2):  # 第 1 行是表头
        values: dict[str, Any] = {}
        try:
            for i, h in enumerate(header):
                if h in cols and i < len(row) and str(row[i]).strip() != "":
                    key, conv = cols[h]
                    try:
                        values[key] = conv(row[i])
                    except ValueError as e:
                        raise ValueError(f"{h} 格式不对：{row[i]}") from e
            obj = schema.model_validate(values)
        except Exception as e:  # noqa: BLE001
            errors.append({"row": line_no, "error": str(e).splitlines()[0][:200]})
            continue
        err = store.check_record(idef, obj)
        if err is None and obj.day > latest_ok:
            err = f"日期 {obj.day} 在未来"
        if err:
            errors.append({"row": line_no, "error": err})
            continue
        by_day[obj.day].append(obj)

    if errors or not by_day:
        rec.status, rec.error_rows = "invalid", errors or [{"row": 0, "error": "没有数据行"}]
        s.flush()
        return rec

    for day, objs in sorted(by_day.items()):
        rec.rows += store.save(s, shop, idef, day, objs, "manual")
        _mark_success(s, parser.item, day, len(objs))
    rec.status = "imported"
    s.flush()
    return rec


def _mark_success(s: Session, item: str, day: date, rows: int) -> None:
    job = s.scalar(select(FetchJob).where(FetchJob.item == item, FetchJob.day == day))
    if job is None:
        job = FetchJob(item=item, day=day, started_at=utcnow())
        s.add(job)
    job.status, job.channel, job.rows, job.error, job.finished_at = "success", "manual", rows, "", utcnow()
    s.flush()
    emit(s, "data.ready", "fetch_job", job.id, {"item": item, "day": day.isoformat(), "channel": "manual"})
